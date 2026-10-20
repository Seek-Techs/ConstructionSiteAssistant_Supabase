
from docx import Document

from helpers import document_text
from observation_model import SiteObservation
from word_report_builder import build_word_report


def test_word_report_contains_observation_details(tmp_path):
    observation = SiteObservation(
        observation_id="OBS-0001",
        report_date="2026-10-15",
        site="Test Drainage Project",
        location="E8, Chainage 0+360",
        contractor="Example Contractor Ltd.",
        category="Quality",
        information_source="Direct observation",
        description="Base slab thickness requires verification.",
        required_action="Verify thickness against the approved specification.",
        responsible_party="Contractor",
        status="Open",
    )

    output_path = tmp_path / "test_report.docx"

    build_word_report(
        [observation],
        output_path,
        db_path=tmp_path / "test_observations.db",
    )

    assert output_path.exists()

    document = Document(output_path)
    text = document_text(document)

    assert "CONSTRUCTION SITE OBSERVATION REPORT" in text
    assert "OBS-0001" in text
    assert "Test Drainage Project" in text
    assert "Direct observation" in text
    assert "Base slab thickness requires verification." in text
    assert "Verify thickness against the approved specification." in text


# ---- Layout and presentation features -------------------------------


import pytest

from helpers import FakeUpload, count_row, make_png_bytes
from storage import (
    get_observation_photos,
    initialize_database,
    save_observation,
    save_project,
    store_uploaded_photo,
)


def _observation(observation_id="OBS-0001", **overrides):
    values = dict(
        observation_id=observation_id,
        report_date="2026-10-20",
        site="Lekki Site",
        location="Chainage 0+100",
        contractor="Example Contractor Ltd.",
        category="Quality",
        information_source="Contractor-reported",
        description="Backfill reported as compacted.",
        required_action="Witness compaction test.",
        responsible_party="Contractor",
        due_date="2026-10-25",
        project_code="TST_01",
    )
    values.update(overrides)
    return SiteObservation(**values)


def test_report_has_identification_block_and_signoff(tmp_path):
    db = tmp_path / "t.db"
    save_project("TST_01", "Test Reclamation", "Lekki", db)
    out = tmp_path / "r.docx"

    build_word_report(
        [_observation()],
        out,
        db_path=db,
        prepared_by="A. Engineer",
        report_reference="TST/2026/010",
    )
    document = Document(out)
    text = document_text(document)

    assert "Project: Test Reclamation (TST_01)" in text
    assert "Project location: Lekki" in text
    assert "Site: Lekki Site" in text
    assert "Report date: 20 Oct 2026" in text
    assert "Report reference: TST/2026/010" in text
    assert "Prepared by | A. Engineer" in text
    assert "Checked by" in text
    # Contractor claims are not presented as personally observed.
    assert "not directly observed by the preparer" in text
    assert "Action due date: 25 Oct 2026" in text


def test_summary_tables_count_selection_only(tmp_path):
    db = tmp_path / "t.db"
    observations = [
        _observation("OBS-0001", category="HSE", status="Open"),
        _observation("OBS-0002", category="HSE", status="In Progress"),
        _observation(
            "OBS-0003",
            category="Quality",
            status="Closed",
            verification_notes="Inspected on site.",
        ),
    ]
    out = tmp_path / "r.docx"
    build_word_report(observations, out, db_path=db)
    document = Document(out)

    categories = count_row(document, "Quality")
    assert categories == {
        "Quality": "1", "HSE": "2", "Traffic": "0",
        "Other": "0", "Total": "3",
    }
    statuses = count_row(document, "Open")
    assert statuses == {
        "Open": "1", "In Progress": "1", "Closed": "1", "Total": "3",
    }


def test_photo_captions_are_numbered_and_fall_back_to_filename(tmp_path):
    db = tmp_path / "t.db"
    save_observation(_observation(), db)
    store_uploaded_photo(
        "OBS-0001",
        FakeUpload("IMG_1.png", make_png_bytes()),
        db,
        caption="Embankment profile, looking north",
    )
    store_uploaded_photo(
        "OBS-0001", FakeUpload("IMG_2.png", make_png_bytes()), db
    )

    stored = get_observation_photos("OBS-0001", db)
    assert [p["caption"] for p in stored] == [
        "Embankment profile, looking north", ""
    ]

    out = tmp_path / "r.docx"
    build_word_report([_observation()], out, db_path=db)
    text = document_text(Document(out))

    assert "Photo 1: Embankment profile, looking north" in text
    assert "Photo 2: IMG_2.png" in text


def test_legacy_photo_table_gets_caption_column(tmp_path):
    import sqlite3

    db = tmp_path / "old.db"
    with sqlite3.connect(db) as connection:
        connection.execute(
            "CREATE TABLE observation_photos (photo_id INTEGER PRIMARY KEY "
            "AUTOINCREMENT, observation_id TEXT NOT NULL, file_path TEXT "
            "NOT NULL, original_filename TEXT NOT NULL)"
        )
        connection.execute(
            "INSERT INTO observation_photos (observation_id, file_path, "
            "original_filename) VALUES ('OBS-0001', 'x.png', 'x.png')"
        )

    initialize_database(db)
    photos = get_observation_photos("OBS-0001", db)

    assert len(photos) == 1
    assert photos[0]["caption"] == ""
    assert photos[0]["original_filename"] == "x.png"


def test_report_without_project_says_so(tmp_path):
    out = tmp_path / "r.docx"
    build_word_report(
        [_observation(project_code="")],
        out,
        db_path=tmp_path / "t.db",
    )
    assert "Project: Not linked to a project" in document_text(Document(out))
