"""End-to-end workflow test using fictional data and a temporary database.

Follows the agreed MVP scenario: register a project, record an observation,
reopen, edit, attach photos, update status, report.
"""

import shutil
import sqlite3

import pytest
from docx import Document

from helpers import (
    FakeUpload,
    count_row,
    document_text,
    make_png_bytes,
)
from observation_model import SiteObservation, validate_observation
from report_builder import build_whatsapp_report
from storage import (
    assign_project_to_observation,
    get_observation_photos,
    get_observations,
    get_projects,
    next_observation_id,
    resolve_photo_path,
    save_observation,
    save_observation_photo,
    save_project,
    store_uploaded_photo,
    update_observation,
    update_observation_status,
)
from word_report_builder import build_word_report


def new_observation(observation_id, project_code="", **overrides):
    values = dict(
        observation_id=observation_id,
        report_date="2026-10-15",
        site="South Drainage",
        location="E8, Chainage 0+360",
        contractor="Sample Contractor",
        category="HSE",
        information_source="Direct observation",
        description="Edge protection not provided along part of excavation.",
        required_action="Provide edge protection; follow-up inspection needed.",
        responsible_party="Contractor",
        due_date="2026-10-20",
        project_code=project_code,
    )
    values.update(overrides)
    return SiteObservation(**values)


def docx_text(path):
    return document_text(Document(path))


def test_full_workflow(tmp_path):
    db = tmp_path / "observations.db"

    # 1-4. Register project, record an observation linked to it.
    save_project("TST_01", "Test Drainage", "Zone A", db)
    obs_id = next_observation_id(db)
    assert obs_id == "OBS-0001"
    obs = new_observation(obs_id, project_code="TST_01")
    assert validate_observation(obs) == []
    save_observation(obs, db)

    # 5-6. "Reopen": a fresh read returns the saved record.
    saved = get_observations(db)[0]
    assert saved == obs

    # 7. Edit without losing identifier, status, notes or project.
    update_observation_status(obs_id, "In Progress", "Work started.", db)
    edited = new_observation(
        obs_id,
        project_code="",  # edit form carries the project; storage ignores it
        description="Edge protection missing along 12 m of excavation.",
    )
    update_observation(edited, db)
    after_edit = get_observations(db)[0]
    assert after_edit.observation_id == obs_id
    assert after_edit.description.startswith("Edge protection missing")
    assert after_edit.status == "In Progress"
    assert after_edit.verification_notes == "Work started."
    assert after_edit.project_code == "TST_01"

    # 8. Attach photos (two valid files, including a duplicate file name).
    store_uploaded_photo(obs_id, FakeUpload("a.png", make_png_bytes()), db)
    store_uploaded_photo(
        obs_id, FakeUpload("a.PNG", make_png_bytes(20, 60)), db
    )
    photos = get_observation_photos(obs_id, db)
    assert len(photos) == 2
    assert all(resolve_photo_path(p, db) for p in photos)

    # 9. Update status/notes. Contractor claim must not close without notes.
    with pytest.raises(ValueError):
        update_observation_status(obs_id, "Closed", "  ", db)
    update_observation_status(
        obs_id, "Closed", "Follow-up inspection 2026-10-19: in place.", db
    )

    # 10-12. Report from saved records, built against this database.
    selected = get_observations(db)
    out = tmp_path / "report.docx"
    build_word_report(selected, out, db_path=db)
    document = Document(out)
    text = docx_text(out)

    assert "TST_01" in text and "Test Drainage" in text
    assert count_row(document, "Quality")["Total"] == "1"
    assert count_row(document, "Open")["Closed"] == "1"
    assert "selected for this report" in text  # scope statement
    assert "Action due date: 20 Oct 2026" in text
    assert "Follow-up inspection 2026-10-19" in text
    assert len(document.inline_shapes) == 2
    assert "Page " in text

    # 13. An open, unverified item is not described as completed/verified.
    save_observation(new_observation("OBS-0002", project_code="TST_01"), db)
    out2 = tmp_path / "report2.docx"
    build_word_report(get_observations(db), out2, db_path=db)
    text2 = docx_text(out2)
    assert "Status: Open" in text2
    assert "No photographs attached to this observation." in text2

    # 14. Earlier records and photos remain available.
    assert len(get_observations(db)) == 2
    assert len(get_observation_photos(obs_id, db)) == 2


def test_report_ignores_other_databases(tmp_path):
    """Report must use the database it is given, not a default one."""
    db = tmp_path / "a.db"
    other = tmp_path / "b.db"
    save_project("P1", "Project One", "Loc", db)
    save_project("P1", "Different Name", "Elsewhere", other)
    save_observation(new_observation("OBS-0001", "P1"), db)
    out = tmp_path / "r.docx"
    build_word_report(get_observations(db), out, db_path=db)
    text = docx_text(out)
    assert "Project One" in text
    assert "Different Name" not in text


def test_photo_requires_existing_observation(tmp_path):
    db = tmp_path / "t.db"
    with pytest.raises(ValueError, match="Observation not found"):
        save_observation_photo("OBS-9999", "x.png", "x.png", db)
    with pytest.raises(ValueError):
        store_uploaded_photo("OBS-9999", FakeUpload("x.png", make_png_bytes()), db)
    # No orphan file or record is left behind.
    assert not list((tmp_path / "photo_evidence").rglob("*.png")) if (
        tmp_path / "photo_evidence"
    ).exists() else True


def test_photos_survive_moving_the_project_folder(tmp_path):
    old = tmp_path / "old"
    old.mkdir()
    db = old / "observations.db"
    save_observation(new_observation("OBS-0001"), db)
    store_uploaded_photo("OBS-0001", FakeUpload("p.png", make_png_bytes()), db)

    new = tmp_path / "moved"
    shutil.move(str(old), str(new))
    new_db = new / "observations.db"

    photo = get_observation_photos("OBS-0001", new_db)[0]
    assert not __import__("pathlib").Path(photo["file_path"]).exists()
    assert resolve_photo_path(photo, new_db) is not None

    out = tmp_path / "r.docx"
    build_word_report(get_observations(new_db), out, db_path=new_db)
    assert len(Document(out).inline_shapes) == 1


def test_bad_or_missing_photo_does_not_break_report(tmp_path):
    db = tmp_path / "t.db"
    save_observation(new_observation("OBS-0001"), db)
    folder = tmp_path / "photo_evidence" / "OBS-0001"
    folder.mkdir(parents=True)
    bad = folder / "bad.png"
    bad.write_bytes(b"not an image")
    save_observation_photo("OBS-0001", str(bad), "bad.png", db)
    save_observation_photo("OBS-0001", str(folder / "gone.png"), "gone.png", db)

    out = tmp_path / "r.docx"
    build_word_report(get_observations(db), out, db_path=db)
    text = docx_text(out)
    assert "Photo could not be embedded: bad.png" in text
    assert "Photo file unavailable: gone.png" in text


def test_unlisted_category_is_not_dropped(tmp_path):
    db = tmp_path / "t.db"
    odd = new_observation("OBS-0001", category="Environmental")
    out = tmp_path / "r.docx"
    build_word_report([odd], out, db_path=db)
    text = docx_text(out)
    assert "Environmental" in text and "OBS-0001" in text
    assert count_row(Document(out), "Quality")["Other"] == "1"
    assert "OBS-0001" in build_whatsapp_report([odd])


def test_assign_project_only_fills_missing_link(tmp_path):
    db = tmp_path / "t.db"
    save_project("P1", "One", "L", db)
    save_project("P2", "Two", "L", db)
    save_observation(new_observation("OBS-0001"), db)

    with pytest.raises(ValueError, match="Project not found"):
        assign_project_to_observation("OBS-0001", "NOPE", db)
    with pytest.raises(ValueError, match="Observation not found"):
        assign_project_to_observation("OBS-0404", "P1", db)

    assign_project_to_observation("OBS-0001", "P1", db)
    assert get_observations(db)[0].project_code == "P1"

    with pytest.raises(ValueError, match="already linked"):
        assign_project_to_observation("OBS-0001", "P2", db)
    assert get_observations(db)[0].project_code == "P1"


def test_legacy_database_is_migrated_without_data_loss(tmp_path):
    db = tmp_path / "old.db"
    with sqlite3.connect(db) as c:
        c.execute(
            """CREATE TABLE observations (
            observation_id TEXT PRIMARY KEY, report_date TEXT NOT NULL,
            site TEXT NOT NULL, location TEXT NOT NULL,
            contractor TEXT NOT NULL, category TEXT NOT NULL,
            information_source TEXT NOT NULL, description TEXT NOT NULL,
            required_action TEXT NOT NULL, responsible_party TEXT NOT NULL,
            status TEXT NOT NULL, verification_notes TEXT NOT NULL)"""
        )
        c.execute(
            "INSERT INTO observations VALUES ('OBS-0001','2026-01-01','S',"
            "'L','C','HSE','Direct observation','d','a','R','Open','')"
        )
    rows = get_observations(db)
    assert len(rows) == 1
    assert rows[0].project_code == "" and rows[0].due_date == ""
    assert get_projects(db) == []
    assert get_observation_photos("OBS-0001", db) == []
